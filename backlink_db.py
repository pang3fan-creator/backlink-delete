#!/usr/bin/env python3
"""
backlink_db.py - 外链数据库管理工具

功能：
- export: 数据库 → Excel（安全合并）
- import: Excel → 数据库（安全模式）
- add-submission: 手动添加提交记录（补救用，正常流程用 backlink_submit.py）
- stats: 查看统计

用法：
python3 backlink_db.py export
python3 backlink_db.py import
python3 backlink_db.py stats
python3 backlink_db.py add-submission <site_id> <project> <status> [选项]
"""

import sqlite3
from datetime import datetime
import sys
import os

from backlink_common import DB_PATH, EXCEL_PATH, LOG_FILE, VALID_STATUSES, migrate_database


def load_openpyxl():
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
    except ModuleNotFoundError:
        print('❌ 缺少 openpyxl，Excel 导入/导出不可用。请运行: python3 -m pip install openpyxl')
        sys.exit(1)

    fills = {
        "green": PatternFill(start_color='E2EFDA', end_color='E2EFDA', fill_type='solid'),
        "yellow": PatternFill(start_color='FFF2CC', end_color='FFF2CC', fill_type='solid'),
        "red": PatternFill(start_color='FCE4EC', end_color='FCE4EC', fill_type='solid'),
        "blue": PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid'),
    }
    return openpyxl, Alignment, Font, fills


def now():
    return datetime.now().strftime('%Y-%m-%d %H:%M:%S')


def validate_status(status):
    """验证状态是否合法"""
    if status not in VALID_STATUSES:
        print(f'❌ 非法状态: "{status}"')
        print(f'   允许的状态: {", ".join(VALID_STATUSES)}')
        return False
    return True


def log_to_file(site_url: str, project: str, submit_url: str, name: str, email: str,
                comment: str, result: str, comment_id: str = None, target_url: str = ""):
    """追加一条提交记录到日志文件"""
    LOG_FILE.parent.mkdir(parents=True, exist_ok=True)
    
    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    
    entry = f"""=== {timestamp} ===
站点：{site_url}
项目：{project}
姓名：{name}
邮箱：{email}
网址：{submit_url}
文章：{target_url}
评论：{comment}
结果：{result}"""
    
    if comment_id:
        entry += f"，评论ID: {comment_id}"
    
    entry += "\n\n"
    
    with open(LOG_FILE, "a", encoding="utf-8") as f:
        f.write(entry)


def export():
    """
    数据库 → Excel（安全合并模式）
    
    规则：
    - Excel 有值 → 保留，不动
    - Excel 无值、数据库有值 → 写入
    - 新站点 → 添加新行
    """
    openpyxl, Alignment, Font, fills = load_openpyxl()
    migrate_database(DB_PATH)
    conn = sqlite3.connect(DB_PATH)
    cursor = conn.cursor()

    # 获取所有项目
    cursor.execute('SELECT DISTINCT project_name FROM submissions ORDER BY project_name')
    projects = [r[0] for r in cursor.fetchall()]

    # 获取所有站点
    cursor.execute('''SELECT id, site_url, site_type, weight
        FROM sites ORDER BY weight DESC NULLS LAST''')
    sites = cursor.fetchall()

    # 获取所有提交记录
    cursor.execute('SELECT site_id, project_name, status FROM submissions')
    db_subs = {}
    for sid, proj, status in cursor.fetchall():
        db_subs[(sid, proj)] = status
    conn.close()

    # 读取现有 Excel 数据（如果存在）
    excel_subs = {}  # (site_url, project) -> status
    excel_site_info = {}  # site_url -> {type, weight}
    
    if os.path.exists(EXCEL_PATH):
        try:
            wb_old = openpyxl.load_workbook(EXCEL_PATH)
            ws_old = wb_old.active
            old_headers = [cell.value for cell in ws_old[1]]
            old_projects = [h for h in old_headers[3:] if h]
            
            for row in ws_old.iter_rows(min_row=2, values_only=True):
                site_url = str(row[0]).strip().lower() if row[0] else ''
                if not site_url:
                    continue
                
                # 保存站点信息
                excel_site_info[site_url] = {
                    'type': row[1],
                    'weight': row[2]
                }
                
                # 保存提交状态
                for proj_idx, proj in enumerate(old_projects):
                    col = 3 + proj_idx
                    if col < len(row) and row[col]:
                        excel_subs[(site_url, proj)] = str(row[col]).strip()
            
            wb_old.close()
            print(f'📖 读取现有 Excel: {len(excel_site_info)} 个站点, {len(excel_subs)} 条提交记录')
        except Exception as e:
            print(f'⚠️ 读取 Excel 失败: {e}，将创建新文件')
            excel_subs = {}
            excel_site_info = {}

    # 创建新 Excel
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '提交总表'

    # 建立 site_url -> site_id 映射
    site_url_to_id = {site[1].lower(): site[0] for site in sites}
    
    # 建立反向映射
    site_id_to_url = {site[0]: site[1].lower() for site in sites}

    headers = ['提交地址', '类型', '权重'] + list(projects)
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = fills["blue"]
        cell.alignment = Alignment(horizontal='center')

    merged_count = 0
    preserved_count = 0

    for row_idx, site in enumerate(sites, 2):
        sid, url, stype, weight = site
        site_url = url.lower()
        
        ws.cell(row=row_idx, column=1, value=url or '')
        ws.cell(row=row_idx, column=2, value=stype or '')
        ws.cell(row=row_idx, column=3, value=str(weight or ''))
        
        for proj_idx, proj in enumerate(projects):
            col = 4 + proj_idx
            
            # 优先使用 Excel 中的值（主人手动填的优先）
            excel_status = excel_subs.get((site_url, proj), '')
            db_status = db_subs.get((sid, proj), '')
            
            if excel_status:
                # Excel 有值，保留
                status = excel_status
                preserved_count += 1
            elif db_status:
                # Excel 无值，数据库有值，写入
                status = db_status
                merged_count += 1
            else:
                status = ''
            
            cell = ws.cell(row=row_idx, column=col, value=status)
            cell.alignment = Alignment(horizontal='center')
            if status == '已提交':
                cell.fill = fills["green"]
            elif status in ('需付费', '需登录', '需验证码', '待确认'):
                cell.fill = fills["yellow"]
            elif status in ('失败', '已失效', '跳过'):
                cell.fill = fills["red"]

    ws.column_dimensions['A'].width = 50
    ws.column_dimensions['B'].width = 15
    ws.column_dimensions['C'].width = 8
    ws.column_dimensions['D'].width = 10

    wb.save(EXCEL_PATH)
    print(f'✅ 导出成功: {EXCEL_PATH}')
    print(f'   站点: {len(sites)} 个, 项目: {len(projects)} 个')
    print(f'   保留 Excel 数据: {preserved_count} 条')
    print(f'   合并数据库数据: {merged_count} 条')


def import_from_excel():
    """
    Excel → 数据库（安全模式）
    
    规则：
    - Excel 有值 → UPSERT 进数据库
    - Excel 无值 → 跳过，不删除数据库记录
    """
    if not os.path.exists(EXCEL_PATH):
        print(f'❌ Excel 文件不存在: {EXCEL_PATH}')
        return

    openpyxl, _, _, _ = load_openpyxl()
    migrate_database(DB_PATH)
    conn = sqlite3.connect(DB_PATH)
    cursor = conn.cursor()
    cursor.execute('PRAGMA foreign_keys = ON')

    wb = openpyxl.load_workbook(EXCEL_PATH)
    ws = wb.active

    # 第一行是表头
    headers = [cell.value for cell in ws[1]]
    projects = [h for h in headers[3:] if h]  # 第4列开始是项目

    # 从第2行开始
    updated = 0
    added = 0
    skipped = 0
    invalid = 0
    now_ts = now()

    for row in ws.iter_rows(min_row=2, values_only=True):
        site_url = str(row[0]).strip().lower() if row[0] else ''
        if not site_url:
            continue
        
        # 确保站点存在
        cursor.execute('INSERT OR IGNORE INTO sites (site_url, created_at) VALUES (?, ?)',
                      (site_url, now_ts))
        
        # 更新站点信息
        if row[1]:
            cursor.execute('UPDATE sites SET site_type = ? WHERE site_url = ?', (row[1], site_url))
        if row[2]:
            try:
                weight = int(row[2]) if row[2] else None
                cursor.execute('UPDATE sites SET weight = ? WHERE site_url = ?', (weight, site_url))
            except:
                pass
        
        cursor.execute('SELECT id FROM sites WHERE site_url = ?', (site_url,))
        site_id = cursor.fetchone()[0]
        
        # 处理每个项目的提交状态
        for proj_idx, proj in enumerate(projects):
            col = 4 + proj_idx
            status = str(row[col]).strip() if col < len(row) and row[col] else ''
            
            if status:
                if not validate_status(status):
                    invalid += 1
                    continue
                # Excel 有值，UPSERT 进数据库
                cursor.execute('''
                    INSERT INTO submissions (site_id, project_name, status, updated_at) 
                    VALUES (?, ?, ?, ?)
                    ON CONFLICT(site_id, project_name) DO UPDATE SET status=?, updated_at=?
                ''', (site_id, proj, status, now_ts, status, now_ts))
                updated += 1
            # Excel 无值，跳过，不删除数据库记录

    conn.commit()

    cursor.execute('SELECT COUNT(*) FROM sites')
    sites_count = cursor.fetchone()[0]
    cursor.execute('SELECT COUNT(*) FROM submissions')
    subs_count = cursor.fetchone()[0]
    conn.close()

    print(f'✅ 导入成功!')
    print(f'   更新/新增提交记录: {updated} 条')
    if invalid:
        print(f'   ⚠️ 跳过非法状态: {invalid} 条（仅允许: {", ".join(VALID_STATUSES)}）')
    print(f'   数据库站点: {sites_count} 个, 提交记录: {subs_count} 条')
    print(f'   💡 空单元格已跳过，未删除数据库记录')


def add_submission():
    """添加/更新提交记录（数据库 + 日志）"""
    if len(sys.argv) < 5:
        print('用法: python3 backlink_db.py add-submission <site_id> <project> <status> [选项]')
        print(f'状态可选: {", ".join(VALID_STATUSES)}')
        print()
        print('选项:')
        print('  --notes "备注"          添加备注')
        print('  --comment "评论内容"    记录日志（博客评论必填）')
        print('  --comment-id "ID"       评论ID（成功时填写）')
        print('  --url "文章URL"         目标文章网址（博客评论必填）')
        print('  --submit-url "网址"     Website 字段填写的外链 URL')
        print('  --reason "原因"         结构化成功/失败原因')
        print('  --name "姓名"           填写的姓名（默认: Stefan M.）')
        print('  --email "邮箱"          填写的邮箱（默认: pang3fan@gmail.com）')
        sys.exit(1)
    
    site_id = sys.argv[2]
    project = sys.argv[3]
    status = sys.argv[4]
    
    # 提取可选参数
    notes = ''
    comment = ''
    comment_id = ''
    target_url = ''
    submit_url = ''
    result_reason = ''
    name = 'Stefan M.'
    email = 'pang3fan@gmail.com'
    
    for i in range(5, len(sys.argv)):
        if sys.argv[i] == '--notes' and i + 1 < len(sys.argv):
            notes = sys.argv[i + 1]
        elif sys.argv[i] == '--comment' and i + 1 < len(sys.argv):
            comment = sys.argv[i + 1]
        elif sys.argv[i] == '--comment-id' and i + 1 < len(sys.argv):
            comment_id = sys.argv[i + 1]
        elif sys.argv[i] == '--url' and i + 1 < len(sys.argv):
            target_url = sys.argv[i + 1]
        elif sys.argv[i] == '--submit-url' and i + 1 < len(sys.argv):
            submit_url = sys.argv[i + 1]
        elif sys.argv[i] == '--reason' and i + 1 < len(sys.argv):
            result_reason = sys.argv[i + 1]
        elif sys.argv[i] == '--name' and i + 1 < len(sys.argv):
            name = sys.argv[i + 1]
        elif sys.argv[i] == '--email' and i + 1 < len(sys.argv):
            email = sys.argv[i + 1]
    
    # 校验状态
    if not validate_status(status):
        sys.exit(1)
    if not result_reason and notes:
        result_reason = notes
    
    # 校验博客评论必填参数
    if comment and not target_url:
        print('❌ 博客评论提交必须提供 --url 参数')
        sys.exit(1)
    
    migrate_database(DB_PATH)
    conn = sqlite3.connect(DB_PATH)
    cursor = conn.cursor()
    cursor.execute('PRAGMA foreign_keys = ON')
    now_ts = now()
    
    try:
        cursor.execute('''
            INSERT INTO submissions (
                site_id, project_name, status, notes, updated_at,
                submit_url, target_url, comment_text, comment_id, result_reason
            ) 
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            ON CONFLICT(site_id, project_name) DO UPDATE SET
                status=?,
                notes=?,
                updated_at=?,
                submit_url=?,
                target_url=?,
                comment_text=?,
                comment_id=?,
                result_reason=?
        ''', (
            site_id, project, status, notes, now_ts, submit_url, target_url, comment, comment_id, result_reason,
            status, notes, now_ts, submit_url, target_url, comment, comment_id, result_reason
        ))
        conn.commit()
        
        # 查询站点信息用于展示和日志
        cursor.execute('SELECT site_url FROM sites WHERE id = ?', (site_id,))
        site = cursor.fetchone()
        site_url = site[0] if site else ''
        
        print(f'✅ 数据库记录已保存')
        print(f'   站点: {site_url}')
        print(f'   项目: {project}')
        print(f'   状态: {status}')
        if notes:
            print(f'   备注: {notes}')
        if result_reason:
            print(f'   原因: {result_reason}')
        
        # 如果有评论内容，写入日志
        if comment:
            log_to_file(
                site_url=site_url,
                project=project,
                submit_url=submit_url,
                name=name,
                email=email,
                comment=comment,
                result='成功' if status == '已提交' else status,
                comment_id=comment_id,
                target_url=target_url
            )
            print(f'✅ 日志已记录到 {LOG_FILE}')
        
    except sqlite3.IntegrityError as e:
        print(f'❌ 数据库错误: {e}')
        conn.rollback()
        sys.exit(1)
    finally:
        conn.close()


def stats():
    """查看统计信息"""
    migrate_database(DB_PATH)
    conn = sqlite3.connect(DB_PATH)
    cursor = conn.cursor()
    
    cursor.execute('SELECT COUNT(*) FROM sites')
    sites_count = cursor.fetchone()[0]
    
    cursor.execute('SELECT COUNT(*) FROM submissions')
    subs_count = cursor.fetchone()[0]
    
    cursor.execute('''
        SELECT site_type, COUNT(*) 
        FROM sites 
        WHERE site_type IS NOT NULL AND site_type != ''
        GROUP BY site_type
        ORDER BY COUNT(*) DESC
    ''')
    types = cursor.fetchall()
    
    cursor.execute('''
        SELECT project_name, COUNT(*) 
        FROM submissions 
        GROUP BY project_name
    ''')
    projects = cursor.fetchall()
    
    cursor.execute('''
        SELECT status, COUNT(*) 
        FROM submissions 
        GROUP BY status
    ''')
    statuses = cursor.fetchall()
    
    conn.close()
    
    print(f'📊 数据库统计')
    print(f'   站点总数: {sites_count}')
    print(f'   提交记录: {subs_count}')
    print()
    print('提交状态:')
    for status, count in statuses:
        print(f'   {status}: {count}')
    print()
    print('站点类型:')
    for stype, count in types:
        print(f'   {stype}: {count}')
    print()
    print('项目提交:')
    for proj, count in projects:
        print(f'   {proj}: {count}')


if __name__ == '__main__':
    if len(sys.argv) < 2:
        print('用法:')
        print('  python3 backlink_db.py export                    → 数据库 → Excel（安全合并）')
        print('  python3 backlink_db.py import                    → Excel → 数据库（安全模式）')
        print('  python3 backlink_db.py add-submission <id> <project> <status> [选项]')
        print('    → 添加/更新提交记录（数据库 + 日志）')
        print()
        print('  python3 backlink_db.py stats                     → 查看统计')
        sys.exit(0)
    
    cmd = sys.argv[1]
    if cmd == 'export':
        export()
    elif cmd == 'import':
        import_from_excel()
    elif cmd == 'add-submission':
        add_submission()
    elif cmd == 'stats':
        stats()
    else:
        print(f'未知命令: {cmd}')
        print('可用命令: export, import, add-submission, stats')
        sys.exit(1)
